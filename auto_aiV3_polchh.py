import os
import io
import json
import pandas as pd
import streamlit as st
import google.generativeai as genai
from datetime import datetime
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ✅ 讀取環境變數
from dotenv import load_dotenv
load_dotenv()
API_KEY = os.getenv("GEMINI_API_KEY")
if not API_KEY:
    st.error("❌ 找不到環境變數 GEMINI_API_KEY，請先設定後重新執行")
else:
    genai.configure(api_key=API_KEY)

# ✅ 模型與欄位設定
model = genai.GenerativeModel("gemini-2.5-flash")
COLUMN_ORDER = [
    "英文名字＋英文姓氏",
    "中文姓名",
    "現職公司",
    "現職職稱",
    "教育背景(學校)",
    "教育背景(系所)",
    "工作經歷",
    "專業技術/證照",
    "專長領域",
    "審查意見或備注"
]

OCR_PROMPT = """
角色：你是一個精準的資料分析人員與審查人員，專門從複雜文件中提取關鍵資訊。
任務：請你仔細閱讀申請人綜合資料 PDF，它包含CV簡歷、背景補充說明等，資料會同時包含中文、英文或其外文。
工作經歷判斷原則：
1.先判斷申請人現職或曾任的公司是否「本身即為軟體或資訊服務業公司」。這表示公司的主營業務必須是軟體開發、系統整合、資訊服務、雲端服務等與軟體或資訊科技直接相關的產業類別。
2.公司如非屬上述軟體或資訊服務業，判斷的重心為申請人的「職務內容」，須「明確顯示其主要工作涉及軟體開發、軟體工具應用或具備專業軟體技術專長」。


輸出格式要求： 請你只輸出一個 JSON 格式的字串 (JSON string)，其中包含所有我要求提取的欄位及其對應的值。請確保每個欄位名稱與我未來 Google Sheet 中的標題完全一致。
如果某個欄位找不到資訊，請將其值設定為 "N/A" (不可用)。
以下是我要求提取的欄位名稱 (JSON Key) 及其預期的值類型/說明：
[
{
"英文名字＋英文姓氏": "string",
"中文姓名": "string",
"現職公司": "string",
"現職職稱": "string",
“教育背景(學校)": "string，因學歷涵蓋大學、研究所和博士等，如有提供不同階段的學歷證明或說明，選擇最高學歷的學校為主",
“教育背景(系所)": "string",輸出格式為 學位_系所名稱或專業領域名稱”,
"工作經歷": "string，請詳細說明每份工作經驗，格式為：每一筆資料是，較早年/月~較晚年/月_任職企業_任職職稱，每筆資料以年份新舊來排列，從最新年份開始排列到過去較舊年份”。
"專業技術/證照": "string",包括會使用的數位技能或工具，證照的格式為：年度＿證照名稱＿發證機構”,
"專長領域": "string，依據工作經驗進判斷。”
"審查意見或備注": "string，針對「工作經驗」每一筆資料，以100字摘要整理出他在這公司的工作內容與事蹟成就，中文回復。並請根據「申請資格條件」與「工作經歷判斷原則」對申請人的工作經歷進行綜合評估。首先判斷申請人所屬公司是否「本身即為軟體或資訊服務業公司」。若非，則判斷申請人「職務內容」是否「明確顯示其主要工作涉及軟體開發、軟體工具應用或具備專業軟體技術專長」。請詳細闡述評估過程與結果，指出申請人的工作經驗是否符合數位經濟相關產業或專業技術要求，若有不符或需進一步釐清之處，請具體說明。例如，若公司非軟體業，且職務偏向數位工具使用者而非開發者，則應在備註中說明。",
}
]


永遠輸出為 JSON 陣列，不使用巢狀JSON。
"""
# ✅ Streamlit UI
st.title("📂 金卡檔案 OCR → Excel 轉換工具 ")

uploaded_files = st.file_uploader(
    "請上傳多個 PDF 或影像檔 (PDF, PNG, JPG)", 
    type=["pdf", "png", "jpg", "jpeg"], 
    accept_multiple_files=True
)

if uploaded_files:
    all_records = []

    with st.spinner("🚀 正在處理所有檔案，請稍候..."):
        for file in uploaded_files:
            st.write(f"🔍 處理檔案：**{file.name}** ...")
            try:
                uploaded = genai.upload_file(file, mime_type=file.type)
                response = model.generate_content([OCR_PROMPT, uploaded])
                ocr_text = response.text.strip()
                ocr_text = re.sub(r"^```json\s*", "", ocr_text)
                ocr_text = re.sub(r"```$", "", ocr_text)
                ocr_text = ocr_text.replace("\ufeff", "").strip()

                st.subheader(f"📜 Gemini 回傳內容 ({file.name})")
                st.code(ocr_text, language="json")

                if not ocr_text:
                    st.error(f"❌ {file.name} OCR 失敗：Gemini API 沒有回傳內容")
                    continue

                try:
                    records = json.loads(ocr_text)
                    if isinstance(records, list):
                        for r in records:
                            r["來源檔案"] = file.name
                        all_records.extend(records)
                        st.success(f"✅ {file.name} 解析完成，共 {len(records)} 筆資料")
                    else:
                        st.warning(f"⚠️ {file.name} 的回傳不是 JSON 陣列格式")
                except json.JSONDecodeError:
                    st.error(f"❌ {file.name} OCR 回傳不是有效的 JSON，請檢查上方內容")
            except Exception as e:
                st.error(f"❌ {file.name} 處理失敗：{e}")

    # 匯總結果轉 Excel
    # 匯總結果轉 Excel
    if all_records:
        df = pd.DataFrame(all_records)
        df = df[[col for col in COLUMN_ORDER if col in df.columns] + ["來源檔案"]]
        st.success(f"🎉 所有檔案完成，共解析 {len(df)} 筆資料")

        # 產生 Excel 檔案（先寫入）
        buffer = io.BytesIO()
        df.to_excel(buffer, index=False, engine="openpyxl")
        buffer.seek(0)

        # 🔧 使用 openpyxl 進行美化：固定欄位寬度、標題顏色、字型與換行
        wb = load_workbook(buffer)
        ws = wb.active
        default_font = Font(size=14)

        # 1️⃣ 標題列美化
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=14)
            cell.fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 2️⃣ 欄位寬度設定
        width_map = {
               "英文名字＋英文姓氏": 27,
                "中文姓名": 14,
                "現職公司": 28,
                "現職職稱": 24,
                "教育背景(學校)": 28,
                "教育背景(系所)": 30,
                "工作經歷": 80,         # 文字多 → 加寬 + wrap_text
                "專業技術/證照": 40,     # 可能多筆
                "專長領域": 28,
                "審查意見或備注": 90
        }

        # 套用欄位寬度
        for idx, col_name in enumerate(df.columns, start=1):
            if col_name in width_map:
                col_letter = ws.cell(row=1, column=idx).column_letter
                ws.column_dimensions[col_letter].width = width_map[col_name]

        # 3️⃣ 設定列高與文字格式（全欄統一）
        for row in range(2, ws.max_row + 1):
            ws.row_dimensions[row].height = 180  # 固定高度，避免換行被擠壓

        # 4️⃣ 全欄套用字體、換行、靠上對齊
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.value:
                    cell.font = default_font
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

        # 儲存回緩衝區
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        

        # 提供下載
        now = datetime.now().strftime("%Y%m%d_%H%M%S")
        st.download_button(
            label="⬇️ 下載匯總 Excel",
            data=buffer,
            file_name=f"OCR結果_{now}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        # ✅ Excel 預覽
        st.subheader("📊 Excel 預覽")
        st.dataframe(df, use_container_width=True)